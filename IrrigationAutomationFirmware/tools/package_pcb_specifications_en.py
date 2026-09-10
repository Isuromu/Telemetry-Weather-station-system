"""Check English document parity and create the requested four-file ZIP."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import re
import shutil
from docx import Document
from pypdf import PdfReader
from build_pcb_specifications_v1 import ROOT, C6, S3, EXP

def normalize(text):
    return re.sub(r'\s+', '', text).replace('\u00ad', '')

items = []
for label, stem, folder, pins, pages in [
    ('01_Soil_Node', 'AmudarIO_Soil_Node_Technical_Specification_v1_0_EN', 'soil', C6, 5),
    ('02_Universal_12V_Controller', 'AmudarIO_Universal_12V_Controller_Technical_Specification_v1_0_EN', 'universal', S3, 7),
]:
    docx = ROOT / 'output' / 'documents' / (stem + '.docx')
    rendered = ROOT / 'tmp' / 'pdfs' / 'pcb_en' / folder / (stem + '.pdf')
    pdf = ROOT / 'output' / 'pdf' / (stem + '.pdf')
    doc = Document(docx)
    reader = PdfReader(rendered)
    text = ''.join(p.extract_text() for p in reader.pages)
    normalized = normalize(text)
    strings = [p.text for p in doc.paragraphs]
    strings += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    strings += [p.text for s in doc.sections for p in s.footer.paragraphs]
    missing = [t for t in strings if normalize(t) not in normalized]
    assert not missing, missing
    assert len(reader.pages) == pages
    assert not re.search(r'[\u0400-\u04FF]', text + ''.join(strings))
    assert not re.search(r'\bdraft\b|\bTBD\b|\bPROPOSED\b', text, re.I)
    rows = [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
    actual_gpio = [r[:4] for r in rows if len(r) == 5 and r[0].isdigit()]
    assert actual_gpio == [r[:4] for r in pins]
    if pins is S3:
        actual_exp = [r[:3] for r in rows if len(r) == 4 and re.fullmatch(r'P[0-2][0-7]', r[0])]
        assert actual_exp == [r[:3] for r in EXP]
    shutil.copy2(rendered, pdf)
    for path in [docx, pdf]:
        items.append((path, label + '_Technical_Specification_EN' + path.suffix))
    print(f'PASS {folder}: {pages} pages, {len(strings)} content checks, {len(pins)} GPIO rows, English only')

downloads = Path('C:/Users/user/Downloads')
assert downloads.is_dir()
archive = downloads / 'AmudarIO_PCB_Specifications_EN_v1_0.zip'
index = 2
while archive.exists():
    archive = downloads / f'AmudarIO_PCB_Specifications_EN_v1_0_{index}.zip'
    index += 1
with ZipFile(archive, 'x', ZIP_DEFLATED, compresslevel=9) as z:
    for path, name in items:
        z.write(path, name)
with ZipFile(archive) as z:
    assert z.testzip() is None
    assert z.namelist() == [name for _, name in items]
    for path, name in items:
        assert hashlib.sha256(z.read(name)).digest() == hashlib.sha256(path.read_bytes()).digest()
    print('ZIP entries:', ', '.join(z.namelist()))
print(f'ARCHIVE: {archive}')
print(f'SIZE: {archive.stat().st_size} bytes')
