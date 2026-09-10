"""Verify final DOCX/PDF text parity and unchanged reviewed pin assignments."""
from pathlib import Path
import re
from docx import Document
from pypdf import PdfReader
from build_pcb_specifications_v1 import ROOT, C6, S3, EXP

def normalize(text):
    return re.sub(r'\s+', '', text).replace('\u00ad', '')

for name, pins, expected_pages in [
    ('AmudarIO_Soil_Node_TZ_v1_0_RU', C6, 5),
    ('AmudarIO_Universal_12V_Controller_TZ_v1_0_RU', S3, 7),
]:
    doc = Document(ROOT / 'output' / 'documents' / (name + '.docx'))
    pdf = PdfReader(ROOT / 'output' / 'pdf' / (name + '.pdf'))
    text = normalize(''.join(p.extract_text() for p in pdf.pages))
    contents = [p.text for p in doc.paragraphs]
    contents += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    missing = [c for c in contents if normalize(c) not in text]
    assert not missing, missing
    assert len(pdf.pages) == expected_pages
    assert not re.search(r'draft|черновик|TBD|PROPOSED', text, re.I)
    rows = [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
    assert all(r in rows for r in pins)
    if pins is S3:
        assert all(r[:4] in rows for r in EXP)
    print(f'{name}: PASS, {len(contents)} text checks, {len(pins)} GPIO rows, {len(pdf.pages)} pages')
