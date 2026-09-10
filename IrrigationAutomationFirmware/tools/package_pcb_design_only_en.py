"""Verify the PCB-only contractor documents and package their matching PDFs."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import re
import shutil
from docx import Document
from pypdf import PdfReader
from build_pcb_specifications_v1 import ROOT, C6, S3, EXP
from build_pcb_design_only_en import C6_P2


def normalize(value):
    return re.sub(r'\s+', '', value).replace('\u00ad', '')


items = []
for label, stem, folder, pins, pages in [
    ('01_Soil_Node', 'AmudarIO_Soil_Node_PCB_Design_Requirements_v1_2_EN', 'soil', C6_P2, 5),
    ('02_Universal_12V_Controller', 'AmudarIO_Universal_12V_PCB_Design_Requirements_v1_2_EN', 'universal', S3, 6),
]:
    docx = ROOT / 'output' / 'documents' / (stem + '.docx')
    rendered = ROOT / 'tmp' / 'pdfs' / 'pcb_mppt_gnss' / folder / (stem + '.pdf')
    pdf = ROOT / 'output' / 'pdf' / (stem + '.pdf')
    doc = Document(docx)
    reader = PdfReader(rendered)
    pdf_text = ''.join(page.extract_text() for page in reader.pages)
    normalized = normalize(pdf_text)
    strings = [paragraph.text for paragraph in doc.paragraphs]
    strings += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    strings += [paragraph.text for section in doc.sections for paragraph in section.footer.paragraphs]
    missing = [value for value in strings if normalize(value) not in normalized]
    assert not missing, missing
    assert len(reader.pages) == pages
    combined = pdf_text + ''.join(strings)
    assert not re.search(r'[\u0400-\u04FF]', combined), 'Non-English text'
    assert not re.search(r'\b(?:draft|TBD|PROPOSED)\b', combined, re.I)
    assert not re.search(r'\b(?:ChirpStack|RX1|RX2|uplink|downlink|ACK|NVS|REG\d+|codec|OTA)\b|\bClass [AC]\b', combined, re.I)
    assert 'Deliver design files only.' in combined
    assert 'physical product testing are outside this assignment' in combined
    assert 'Version 1.2' in combined
    rows = [[cell.text for cell in row.cells] for table in doc.tables for row in table.rows]
    actual_gpio = [row[:4] for row in rows if len(row) == 5 and row[0].isdigit()]
    assert actual_gpio == [row[:4] for row in pins], 'GPIO allocation changed'
    if pins is C6_P2:
        assert [row[:2] for row in pins] == [row[:2] for row in C6]
        changed = {new[0] for old, new in zip(C6, pins) if old[:4] != new[:4]}
        assert changed == {'2', '6', '12', '13'}
        for phrase in ['LTC4121IUD-4.2#PBF', 'MAX-M10S-00B', 'GNSS_SEL', 'PERIPH_PWR_EN', 'V_BCKP', 'TMUX1574', 'fractional-Voc MPPT']:
            assert phrase in combined, phrase
        assert 'LTC4079' not in combined
        assert 'Do not fit GPS' not in combined
    else:
        assert 'MUST use an external MPPT lead-acid charge controller' in combined
    if pins is S3:
        actual_exp = [row[:3] for row in rows if len(row) == 4 and re.fullmatch(r'P[0-2][0-7]', row[0])]
        assert actual_exp == [row[:3] for row in EXP], 'Expander allocation changed'
    pdf.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(rendered, pdf)
    for path in [docx, pdf]:
        items.append((path, label + '_PCB_Design_Requirements_EN_v1_2' + path.suffix))
    print(f'PASS {folder}: {pages} pages, {len(strings)} content checks, {len(pins)} verified GPIO rows')

downloads = Path('C:/Users/user/Downloads')
assert downloads.is_dir()
archive = downloads / 'AmudarIO_PCB_MPPT_GPS_EN_v1_2.zip'
index = 2
while archive.exists():
    archive = downloads / f'AmudarIO_PCB_MPPT_GPS_EN_v1_2_{index}.zip'
    index += 1
with ZipFile(archive, 'x', ZIP_DEFLATED, compresslevel=9) as zipped:
    for path, name in items:
        zipped.write(path, name)
with ZipFile(archive) as zipped:
    assert zipped.testzip() is None
    assert zipped.namelist() == [name for _, name in items]
    for path, name in items:
        assert hashlib.sha256(zipped.read(name)).digest() == hashlib.sha256(path.read_bytes()).digest()
    print('ZIP entries:', ', '.join(zipped.namelist()))
print(f'ARCHIVE: {archive}')
print(f'SIZE: {archive.stat().st_size} bytes')
