"""PCB-design-only contractor scope, retaining the existing electrical baseline."""
from build_pcb_specifications_en import EnglishSpec, gpio, radio, EXP_DESCRIPTIONS, DESCRIPTIONS
from build_pcb_specifications_v1 import C6, S3, EXP

C6_P2 = [r[:4] + [DESCRIPTIONS[r[3]]] for r in C6]
for r in C6_P2:
    if r[0] == '2':
        r[3:] = ['PERIPH_PWR_EN', 'Enable selected soil or GNSS domain and UART mux']
    elif r[0] == '6':
        r[2:] = ['O', 'GNSS_SEL', '0 = soil; 1 = GNSS; gated by GPIO2']
    elif r[0] == '12':
        r[3:] = ['PERIPH_UART_TX', 'UART1 TX to soil or GNSS through mux']
    elif r[0] == '13':
        r[3:] = ['PERIPH_UART_RX', 'UART1 RX from selected soil or GNSS branch']

class DesignSpec(EnglishSpec):
    def __init__(self, stem, title, board, pinout):
        super().__init__(stem,title,board,pinout)
        subtitle = self.doc.paragraphs[1]
        for run in subtitle.runs:
            run.text = run.text.replace('Version 1.0','Version 1.2')
        self.md[2] = self.md[2].replace('Version 1.0','Version 1.2')
        self.doc.sections[0].footer.paragraphs[0].runs[0].text = 'Version 1.2  |  '
        self.doc.core_properties.subject = 'Electrical schematic and PCB design requirements'
        self.doc.core_properties.keywords = 'PCB design, version 1.2, MPPT, GNSS, English'
    def scope(self):
        self.p('Scope: electrical schematic capture and PCB design for the hardware specified below. Deliver design files only. Firmware, command protocols, cloud services, enclosure design, manufacturing, assembly, commissioning and physical product testing are outside this assignment. Board outline, mounting coordinates and external interface requirements are supplied by the customer.')

def service(s,tx,rx,boot,en):
    s.table(['Pin','Board net','Connection'],[
        ['1','GND','Common ground'], ['2','3V3_REF','High-impedance adapter reference only; not a power input'],
        ['3',f'UART0_TX / GPIO{tx}','External adapter RX'], ['4',f'UART0_RX / GPIO{rx}','External adapter TX'],
        ['5','DTR','Input to onboard two-transistor auto-reset circuit'],
        ['6','RTS','Input to onboard two-transistor auto-reset circuit'],
    ],[.65,2.25,4.4])
    s.p(f'External USB-to-TTL interface: 3.3 V logic only. The board uses its own power source. Do not connect adapter 3.3 V, 5 V or VBUS power outputs to the board. Route DTR/RTS through the standard Espressif two-transistor circuit, not directly to the MCU. Provide independent BOOT (GPIO{boot}) and RESET (EN, module pad {en}) buttons or pads. No onboard USB-UART bridge is required.')

def deliverables(s):
    s.p('Deliver the native schematic/PCB project with libraries; schematic PDF; BOM with exact manufacturer part numbers and DNP variants; Gerber and NC drill files; fabrication notes and stackup; pick-and-place and assembly drawings; PCB-only STEP model; ERC/DRC results with justified exceptions; RF impedance geometry and electrical sizing notes. Assembly files are design outputs, not an obligation to manufacture or assemble boards.')

def soil():
    s=DesignSpec('AmudarIO_Soil_Node_PCB_Design_Requirements_v1_2_EN','Soil node PCB design requirements',1,'C6-P2')
    s.scope()
    s.h('1 Hardware composition')
    s.table(['Block','Required implementation'],[
        ['MCU','ESP32-C6-MINI-1-H4, 4 MB flash; pin map C6-P2'],
        ['Radio','SX1262-based DX-LR30-900M22S, 3.3 V; external LoRa antenna via SMA'],
        ['Battery','Removable protected 1S 18650 Li-ion cell, 3.6/3.7 V nominal, 4.2 V charge'],
        ['Holder','MPD BK-18650-PC2 with retainer; PCB includes holder, not the cell'],
        ['Solar MPPT charger','LTC4121IUD-4.2#PBF; autonomous fractional-Voc MPPT, 4.2 V CC/CV'],
        ['GPS and GNSS','u-blox MAX-M10S-00B; switched 3.3 V supply and separate GNSS antenna'],
        ['UART selection','TMUX1574PW; power gated, soil/GNSS selector on UART1'],
        ['System 3.3 V','TPS63900 buck-boost, always on when SYSTEM is ON'],
        ['Probe 5 V','TPS61023 with EN and true output disconnect'],
        ['RS-485','THVD1406 automatic direction; 3.3 V rail switched by TPS22917'],
        ['Battery ADC','Internal ADC; divider switched on its high side by TPS22917'],
        ['Service','UART0 header, DTR/RTS circuit, BOOT and RESET; no separate wake button'],
    ],[1.65,5.65])
    s.p('MPPT and GPS/GNSS are mandatory fitted functions. Do not fit an external RTC, ADS1115, a fuel gauge or continuously lit LEDs. LiFePO4 is not an interchangeable cell option.')
    s.h('2 Power architecture and MPPT charger')
    s.p('A nominal 6 V panel feeds the charger. The protected battery feeds both the 3.3 V converter and 5 V boost through the mechanical SYSTEM switch. SYSTEM OFF disconnects all system loads; the panel, charger, battery and NTC remain connected. Prevent back-power through service and probe connectors.')
    s.p('Use LTC4121-4.2 with its MPPT divider fitted and enabled. It periodically samples panel open-circuit voltage and regulates a programmed fraction of it; this is autonomous fractional-Voc MPPT, not fixed-Vmp regulation or a full power-curve sweep. MPPT must operate with the MCU asleep and SYSTEM OFF. A non-MPPT substitute is not permitted. [1]')
    s.p('Select the divider ratio from the actual panel Vmp/Voc; do not assume a universal percentage. The charger needs at least 4.4 V and battery headroom after the input blocking element. Check hot-panel Vmp, cold-panel Voc, input capacitance and weak-light startup. Program 50-250 mA within cell, panel and thermal limits. Follow the solar reference circuit: place MPPT/RUN sensing on the panel side of reverse blocking to prevent dark battery drain. [1]')
    s.p('Fit a cell-contact NTC with hardware charge inhibition for unsafe temperature and NTC faults. Independent cell protection covers overvoltage, undervoltage, overload and short circuit; add reverse-insertion protection. Size converter peaks and capacitance from the supplied probe, radio and GNSS loads. The charger is not an MCU-controlled power switch.')

    s.page(); s.h('3 Complete ESP32 C6 pinout')
    s.p('GPIO is the signal identifier; Pad is the physical ESP32-C6-MINI-1-H4 module pad, not a bare-chip pin or DevKit header. A = ADC input; I = input; O = output. Logic level: 3.3 V.')
    s.table(['GPIO','Pad','I/O','Net','Function'],C6_P2,[.43,.43,.35,2.35,3.74],9)
    s.p('Supply: pad 3 = 3.3 V; pad 8 = EN/CHIP_PU. GND: pads 1, 2, 11, 14, 36-53. NC: pads 4, 7, 21, 32-35. GPIO10/11 are not exposed. No spare GPIOs remain. Add EN pull-up, RC and supervisor for slow/interrupted supply ramps.')
    s.p('C6-P2 changes GPIO6 from SERVICE_WAKE_N to GNSS_SEL and routes UART1 through a selector. GPIO2 is now PERIPH_PWR_EN; it powers only the branch selected by GPIO6. GPIO1 still independently enables the battery divider. LoRa and UART0 pin numbers are unchanged. No load current flows through GPIOs.')
    s.p('Do not connect GPIOs directly to battery voltage, 5 V or RS-485 A/B. The extra wake button is omitted; BOOT and RESET remain. No strapping or flash pin is repurposed for GNSS.')

    s.page(); s.h('4 Pin restrictions and electrical defaults')
    s.bullets(
        'GPIO4/5/8/9/15 are strapping pins. Do not load GPIO4/5. Pull GPIO8/9 HIGH through 10 kOhm and GPIO15 LOW through 10 kOhm. BOOT pulls GPIO9 to GND; preserve GPIO8=1 and GPIO9=0 for UART download, GPIO9=1 for normal boot.',
        'GPIO6/7 reuse pad-JTAG functions. Do not connect an external JTAG interface to GPIO4-7 in this design. Respect the factory eFuse/JTAG source configuration; no eFuse changes are part of PCB design.',
        'GPIO12/13 default to USB D-/D+. Both UART paths require powered-off isolation from soil AND GNSS, including startup. The mux must be disconnected while either selected rail is invalid. A series resistor alone does not establish isolation.',
        'GPIO16/17 are dedicated UART0. GPIO18-23 can have startup SDIO pulls. Keep LoRa NSS inactive during startup; do not attach power enables to UART0.',
    )
    s.table(['Nets','External reset defaults','Required low-power levels'],[
        ['GPIO1/2','PD about 100 kOhm; LOW','LOW; divider, soil, GNSS and mux OFF'],
        ['GPIO3/7 RXEN/TXEN','PD about 100 kOhm; LOW','LOW'],
        ['GPIO14/21 RESET/NSS','PU 47-100 kOhm; HIGH','HIGH with radio in shutdown'],
        ['GPIO18/20 SCK/MOSI','Keep NSS HIGH','LOW'],
        ['GPIO6 GNSS_SEL','PD 10 kOhm; GPIO2 blocks all loads','LOW; soil selected but unpowered'],
        ['ADC and UART inputs','Voltage limits and no back-power','No floating or powered-off leakage paths'],
    ],[2.15,2.5,2.65])
    s.p('GPIO6 has a reset-time internal weak pull-up. Size its external pull-down for margin; GPIO2 LOW must independently inhibit both domains regardless of GPIO6. All loads remain off while MCU pins are high impedance. Low-power levels are electrical interface requirements, not a software work item.')
    s.h('5 Radio wiring'); radio(s,True)

    s.page(); s.h('6 GPS and GNSS hardware')
    s.p('Fit MAX-M10S for autonomous outdoor position acquisition, including GPS. Coordinates must be electrically accessible to the ESP32 through UART1 for subsequent transmission by the existing radio. No location-update interval or application protocol is assigned to the PCB designer. GNSS and soil sensing are sequential, not simultaneous. [2]')
    s.table(['MAX-M10S pins','Connection'],[
        ['1, 10, 12 / GND','Ground plane'],
        ['7 / V_IO; 8 / VCC','3V3_GNSS_SW from a TPS22917 load switch'],
        ['2 / TXD','Through UART mux to ESP32 GPIO13'],
        ['3 / RXD','Through UART mux from ESP32 GPIO12'],
        ['6 / V_BCKP','Leave open; no backup cell or always-on backup supply'],
        ['11 / RF_IN','50 ohm path to a separate GNSS U.FL antenna connector'],
        ['15 / VIO_SEL','Leave open for 3.3 V I/O'],
        ['9 / RESET_N','Local reset test pad; no pull-up to an always-on rail'],
        ['4, 5, 13, 14, 16, 17, 18','Unused in this UART and external antenna-bias implementation; leave open'],
    ],[2.3,5.0])
    s.p('Switch VCC and V_IO together; remove antenna bias at the same time. Allow at least 100 mA receiver startup current plus the selected antenna load and margin. Respect the V_IO ramp limits and add local decoupling. Cold starts after full shutdown are expected. No GNSS pin may back-power the switched supply. [2, 3]')
    s.p('Use an external active L1 GNSS antenna rated for 3.3 V through a short coax lead. Feed bias from 3V3_GNSS_SW through a current-limited, filtered bias tee and add low-capacitance RF ESD protection. Select gain, current and bias components per the integration guide. Keep GNSS away from LoRa TX and switching nodes; do not share the LoRa SMA. The antenna needs sky visibility. [3]')
    s.h('7 Power and UART selection')
    s.table(['GPIO2','GPIO6','Powered domain','UART1 connection'],[
        ['0','X','Both OFF; mux supply OFF','Disconnected'],
        ['1','0','Soil 5 V and RS-485 3.3 V only','Soil RS-485'],
        ['1','1','GNSS and antenna only','GNSS'],
    ],[.65,.65,3.2,2.8])
    s.p('Hardware decode: SOIL_EN = GPIO2 AND NOT GPIO6; GNSS_EN = GPIO2 AND GPIO6. Apply reset/brownout inhibition. Power the TMUX1574 from a third switched 3.3 V rail enabled by GPIO2 so its standby current is absent in deep sleep. UART0 is never multiplexed.')
    s.p('TMUX1574PW TSSOP: D1 pin 4 to GPIO12; S1A pin 2 to RS-485 DI; S1B pin 3 to GNSS RXD. D2 pin 7 to GPIO13; S2A pin 5 to RS-485 RO; S2B pin 6 to GNSS TXD. SEL pin 1 to GPIO6; VDD pin 16 to switched mux supply; GND pin 8 to GND. Ground unused channels at pins 9-14. [4]')
    s.p('EN pin 15 is active LOW. Hardware must hold it HIGH until the selected supply is valid, and disconnect it before power removal or branch changes. Use break-before-make gating and specified off-state isolation; the mux alone does not protect an unpowered endpoint while its channel is ON. Both domains must be off during selection changes. Add test pads for both UART branches and supply enables. [4]')

    s.page(); s.h('8 Switched probe rail and battery ADC')
    s.p('SOIL_EN enables TPS61023 and the RS-485 supply switch together. Add output discharge and probe short-circuit protection. A/B and both UART paths must not back-power disabled rails. Use a keyed probe connector: 1 = +5V_SOIL_SW, 2 = GND/reference, 3 = RS485_A, 4 = RS485_B. Panel connector: PV+ and PV-.')
    s.p('RS-485 direction is automatic; no DE/RE GPIO is allocated. Provide connector-side SM712-class TVS, surge-rated series-resistor footprints and optional 120 ohm termination/bias. Bias belongs to the switched domain. Select termination and AutoDirection timing provisions for the supplied cable and baud rate, including 4800/9600 baud requirements.')
    s.p('Battery divider starting values: high-side switch, 499 + 499 kOhm upper leg, 249 kOhm lower leg and 100 nF ADC-to-GND. GPIO1 controls the switch; GPIO0 is ADC1_CH0. Check divider ratio, RC settling, resistor tolerances, leakage and ADC voltage limits. Include a voltage test pad; no calibration code is required.')
    s.h('9 Service connector'); service(s,16,17,9,8)
    s.h('10 PCB layout and design outputs')
    s.p('Use two-layer FR-4 with a nearly continuous GND plane. Four layers require routing/EMC justification. Route ANT-to-SMA as a short 50 ohm transmission line using fabricator-confirmed geometry, not a series 50 ohm resistor. Keep switching loops/inductors away from ADC/RF, respect module antenna keepouts and place protection beside connectors.')
    s.p('Sleep target: complete board <=25 uA; mandatory inactive-state ceiling <1 mA. Conditions: timer deep sleep, radio shutdown, soil/GNSS/antenna/divider/mux OFF, adapter detached and dark panel. Acquisition, sensing, radio activity and charging are active states. Budget all leakage across temperature, including the MPPT charger and isolation; <=25 uA is not yet a validated worst-case result. Do not delete MPPT or GNSS to meet the target. Fit a current-measurement link; physical testing is outside this assignment.')
    deliverables(s)
    s.p('Before design release confirm cell/NTC, panel Voc/Vmp/Isc, probe load, radio revision/band, GNSS antenna, board outline/connector locations and environmental ratings. Do not invent missing load ratings.')
    s.h('11 Component design references')
    for ref in [
        '[1] ADI LTC4121-4.2: analog.com/media/en/technical-documentation/data-sheets/4121fbc.pdf',
        '[2] u-blox MAX-M10S: content.u-blox.com/sites/default/files/MAX-M10S_DataSheet_UBX-20035208.pdf',
        '[3] MAX-M10S integration: content.u-blox.com/sites/default/files/MAX-M10S_IntegrationManual_UBX-20053088.pdf',
        '[4] TI TMUX1574: ti.com/lit/ds/symlink/tmux1574.pdf',
    ]:
        s.p(ref)
        for run in s.doc.paragraphs[-1].runs:
            from docx.shared import Pt
            run.font.size = Pt(8)
    s.save()

def universal():
    s=DesignSpec('AmudarIO_Universal_12V_PCB_Design_Requirements_v1_2_EN','Universal 12 V PCB design requirements',2,'S3-P2')
    s.scope()
    s.h('1 Hardware composition')
    s.table(['Block','Required implementation'],[
        ['MCU','ESP32-S3-WROOM-1-N8, 8 MB flash, no PSRAM; pin map S3-P2'],
        ['Radio','DX-LR30 / SX1262, 3.3 V; external LoRa antenna via SMA'],
        ['Valve outputs','2 x DRV8874, PH/EN mode; separate high-side gates, current sense and nFAULT'],
        ['RS-485','4 protected selectable branches, one active, automatic direction'],
        ['Converters','TPS62933 for always-on 3.3 V and a separate enabled 5 V rail'],
        ['Auxiliary power','Protected switched 5 V output and protected VBAT_SW output'],
        ['I/O','TCA6424A RGJ at 0x22; local switched 3.3 V I2C; 4 protected state inputs'],
        ['Service','UART0 header, DTR/RTS circuit and BOOT/RESET access'],
    ],[1.65,5.65])
    s.p('Use one bare PCB with documented DNP variants for unused H-bridges, RS-485 branches and associated isolated power. No application-mode or software implementation is required from the PCB designer.')
    s.h('2 Input and power protection')
    s.p('Input is nominal 12 VDC from a lead-acid battery or an external certified isolated 230 VAC to 12 VDC supply. No mains voltage may enter the PCB. Solar installations MUST use an external MPPT lead-acid charge controller; PWM-only or fixed-voltage solar chargers are not acceptable. MPPT, battery-specific charge/float settings and temperature compensation remain external; no solar panel connects directly to this PCB. Check the 10.5-14.8 V input range against the chosen charger and source limits.')
    s.p('Include input fuse, reverse-polarity protection, transient/overvoltage protection and hardware undervoltage supervision with hysteresis. TPS37-Q1 is the supervisor baseline; it requires a separately rated load-disconnect stage. Expose POWER_FAULT_N. Provide separate resistor/configuration variants for battery and external-PSU supply. The discussed 11.5 V level is not a fixed whole-board cutoff; use customer-approved warning/inhibit, cutoff and recovery thresholds.')
    s.p('Hardware cutoff must not occur before the customer-specified reserve for a closing pulse is exhausted. Size switches, fuse, copper and bulk capacitance using valve pulse and auxiliary-load data. The design must keep H-bridges off during reset, startup, brownout and service-connector activity.')
    s.h('3 Low-power hardware allocation')
    s.p('The minimal low-power assembly has a design goal of <=50 uA at the 12 V input with MCU asleep, radio in shutdown and external domains OFF. This is not a guaranteed current for every populated variant. Calculate leakage by BOM variant including converter, expander, supervisor, pulls, switches and isolated power. Fit a removable measurement link and no continuously lit power LEDs. Physical tests are outside this assignment.')

    s.page(); s.h('4 Complete ESP32 S3 pinout')
    s.p('GPIO is the signal number; Pad refers to ESP32-S3-WROOM-1-N8. A = ADC input; I = input; O = output; OD = open-drain. Logic level is 3.3 V. Do not substitute classic ESP32 pin numbers.')
    gpio(s,S3[:22])
    s.h('Module power and unavailable pins')
    s.p('Pad 2: 3.3 V; pad 3: EN/CHIP_PU with RC and supervisor. Pads 1, 40 and exposed pad 41: GND. GPIO22-34 are not available as external module GPIOs; do not use internal flash connections. Clamp and scale ADC inputs; no GPIO may connect directly to 5/12/24 V, a power MOSFET gate or RS-485 A/B.')

    s.page(); s.h('4 Complete ESP32 S3 pinout continued'); gpio(s,S3[22:])
    s.h('5 Restricted pins')
    s.bullets(
        'GPIO0/3/45/46 are straps. Pull GPIO0/3 HIGH and GPIO45/46 LOW through 10 kOhm. GPIO0 is BOOT only. No field loads on straps; respect factory eFuse configuration.',
        'Reserve GPIO39-42 for JTAG and GPIO19/20 for USB/debug only. UART0 GPIO43/44 must not control loads; startup output activity is expected.',
        'GPIO35/36/37 require the specified no-PSRAM N8. Octal-PSRAM N8R8 is not interchangeable. R16V uses 1.8 V on GPIO47/48. Any module substitution requires an electrical review.',
        'GPIO15/16 are occupied by LoRa; no 32.768 kHz crystal. GPIO37/48 are not RTC wake pins; do not describe the expander interrupt as a deep-sleep wake source.',
    )
    s.h('6 Radio wiring'); radio(s,False)

    s.page(); s.h('7 TCA6424A expander pinout')
    s.p('32-pin RGJ: VCCI 31 and VCCP 27 to always-on 3.3 V; GND 25 and exposed pad to GND; ADDR 26 to GND for address 0x22. SCL 29 to GPIO7, SDA 30 to GPIO8, INT_N 32 to GPIO37. RESET_N 28 connects to the board reset/watchdog network. Expander port P13 is not MCU GPIO13.')
    s.table(['Port','Pin','Net','Function'],[r[:3]+[d] for r,d in zip(EXP,EXP_DESCRIPTIONS)],[.48,.48,2.65,3.69],9)
    s.p('P00-P17 require external pull-downs of about 100 kOhm; their required inactive state is LOW. P20-P27 are protected inputs with safe pull-ups. At reset all ports are inputs, but the output latches contain ones. Hardware interlocks must prevent load activation while the expander is being configured. Do not rely on expander state alone to energize a bridge.')
    s.p('The reset/watchdog circuit must remove stale output permissions even when an internal MCU reset does not pull EN LOW. GPIO37 cannot wake deep sleep; the four field inputs are state/contact inputs, not guaranteed lossless pulse counters. Protect fault inputs from back-power through unpowered drivers.')
    s.p('GPIO startup glitches must not reach enabled power loads. LoRa NSS/RESET_N: pull HIGH with 47-100 kOhm; RXEN/TXEN: pull LOW with about 100 kOhm. Internal I2C SCL/SDA pull-ups start at 4.7 kOhm. Check all pulls for noise margin and current consumption.')

    s.page(); s.h('8 H bridges and controlled outputs')
    s.table(['Connection','Valve 1','Valve 2'],[
        ['PH polarity','GPIO5','GPIO38'], ['EN pulse input','GPIO6','GPIO47'],
        ['High-side power request','GPIO35','GPIO36'], ['nSLEEP permission','TCA6424A P14','TCA6424A P15'],
        ['nFAULT','TCA6424A P24','TCA6424A P25'], ['IPROPI current ADC','GPIO2 / ADC1_CH1','GPIO4 / ADC1_CH3'],
    ],[2.7,2.3,2.3])
    s.p('Strap both DRV8874 devices for PH/EN operation. Each channel has a separate high-side supply switch, local bulk capacitance, coil-energy suppression, hardware current limiting and fault feedback. PH, EN, nSLEEP and power requests default LOW. Fit about 100 kOhm pull-downs on direct control lines. Scale IPROPI to the ADC range.')
    s.p('Bridge permission requires the direct MCU request, OUTPUT_ARM, valid input power and reset/watchdog permission. nSLEEP may be released only for a powered permitted channel. Both pulse paths must not be enabled simultaneously. Include a bounded-on-time protection mechanism using the customer-supplied maximum pulse duration; do not assign software development to the PCB designer.')
    s.p('Outputs must support opposite-polarity pulses for two-wire latching coils. Power removal alone is not a closing pulse. Design for the supplied pulse current, duration, minimum closing voltage, cable drop and repetition limits. No valve motion or hydraulic control algorithm is part of this assignment.')
    s.h('9 Auxiliary power and battery ADC')
    s.bullets(
        'P07 enables the 5 V buck; P10 independently switches AUX 5 V. P11 switches AUX VBAT. Label the latter VBAT_SW: it follows battery voltage, not regulated 12.0 V.',
        'Provide high-side switching, output overload/short protection, reverse-current blocking as needed, default-OFF control and connector protection. Size copper, switches and connectors for declared loads.',
        'P13 controls the battery-divider high-side switch; GPIO1 is the ADC input. Include RC filtering, ADC voltage protection and a test pad. Calculate divider tolerance, leakage and settling time; no calibration routine is required.',
    )
    s.h('10 Local I2C and field I O')
    s.p('P12 gates the local 3.3 V I2C connector power AND SDA/SCL signal paths. Connector nets: 3V3_SW, GND, SDA, SCL. Keep TCA6424A and internal pulls always powered; an unpowered external device must not clamp the internal bus. This port is for short local wiring.')
    s.p('Protect and condition four field state inputs; dry-contact or 12/24 V implementation is fixed by the customer-selected assembly. P17 drives only an isolated low-voltage equipment-control interface. It is not a safety relay or a mains output. External equipment safety circuits remain outside this PCB; preserve the specified isolation and default-inactive interface.')

    s.page(); s.h('11 RS485 branches and service header')
    s.p('UART1 TX=GPIO18, RX=GPIO21. Use a break-before-make selector with one active branch; never tie push-pull receiver outputs together. SEL1:SEL0 = 00/01/10/11 selects ports 1/2/3/4. P02=0 disconnects all; P02 is not DE/RE. Hardware-decode P03-P06 power requests so only the selected branch is enabled.')
    s.p('Automatic direction is mandatory. Baseline family: MAX22025F/MAX22026F with separate cable-side supply; confirm the exact variant and electrical timing against the supplied baud rates/cable. Provide PWR, RETURN, A, B and SHIELD per port, TVS, current protection and optional 120 ohm termination/bias. Fix 5 V/VBAT selection by assembly, never by a software-selectable voltage switch.')
    s.p('Ports 1/2: pressure-sensor power selected from supplied ratings. Port 3: TUF-2000M, VBAT supply, reference load 8-36 V and about 50 mA. Port 4: spare field instrument. Protocols and register maps are not required for PCB design. Provide disconnectable transceiver and isolated-power domains with no TX/RX, A/B or fault-line back-power.')
    s.p('Full isolation requires isolated data AND field-device power/return. A shared VBAT/5 V ground is signal-only isolation. State the isolation type per BOM variant; size isolated DC/DC for the supplied load and startup current. Do not connect the VFD cable-side return to logic ground. Use customer-specified isolation ratings, clearances and shield termination.')
    service(s,43,44,0,3)
    s.h('12 PCB layout and design outputs')
    s.p('Maximum-feature board: four-layer FR-4 with continuous reference planes except at required isolation boundaries. Keep H-bridge and converter loops short; separate high-current returns from ADC/RF paths. ANT-to-SMA is a short calculated 50 ohm transmission line. Respect module keepouts and place protection at connectors. A two-layer alternative requires customer-approved routing justification.')
    deliverables(s)
    s.p('Customer inputs before design release: valve pulse/load ratings; field-device power and cable data; auxiliary output currents; isolation and DC/DC requirements per port; supply/UVLO thresholds; radio revision and RF band; board outline, mounting/connector coordinates and environmental ratings. Missing values require customer clarification, not assumptions. No enclosure, application code or assembled-board test report is required.')
    s.save()

if __name__=='__main__':
    soil()
    universal()
